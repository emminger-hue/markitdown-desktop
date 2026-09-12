"""Regenerate the sample files used by the test-suite.

Run: python tests/fixtures/make_fixtures.py
All files are tiny and deterministic enough to be committed.
"""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

HERE = Path(__file__).resolve().parent
TEXT_LINES = [
    "Rechnung Nr. 4711 vom 12. März 2026",
    "Straße, Größe, Übung, Äpfel – 1.234,56 €",
    "The quick brown fox jumps over the lazy dog.",
]
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"


def make_png(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (1000, 220), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT, 34)
    for i, line in enumerate(TEXT_LINES):
        draw.text((20, 20 + i * 60), line, fill="black", font=font)
    img.save(path)


def make_text_pdf(path: Path) -> None:
    from PySide6.QtGui import QFont, QPageSize, QPainter, QPdfWriter
    from PySide6.QtWidgets import QApplication

    QApplication.instance() or QApplication([])
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(96)
    painter = QPainter(writer)
    painter.setFont(QFont("DejaVu Sans", 14))
    y = 80
    for line in ["Textseite eins", *TEXT_LINES]:
        painter.drawText(60, y, line)
        y += 30
    painter.end()


def make_scan_pdf(path: Path, png: Path) -> None:
    from PySide6.QtGui import QImage, QPageSize, QPainter, QPdfWriter
    from PySide6.QtWidgets import QApplication

    QApplication.instance() or QApplication([])
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(96)
    painter = QPainter(writer)
    painter.drawImage(40, 40, QImage(str(png)))
    painter.end()


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Beispieldokument", level=1)
    for line in TEXT_LINES:
        document.add_paragraph(line)
    document.save(str(path))


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Beispielfolie"
    slide.placeholders[1].text = "\n".join(TEXT_LINES)
    prs.save(str(path))


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Daten"
    ws.append(["Position", "Betrag"])
    ws.append(["Äpfel", 12.5])
    ws.append(["Übung", 7])
    wb.save(str(path))


def main() -> None:
    png = HERE / "sample.png"
    make_png(png)
    make_text_pdf(HERE / "sample.pdf")
    make_scan_pdf(HERE / "scan.pdf", png)
    make_docx(HERE / "sample.docx")
    make_pptx(HERE / "sample.pptx")
    make_xlsx(HERE / "sample.xlsx")
    (HERE / "sample.html").write_text(
        "<html><body><h1>Beispiel</h1><p>Straße, Größe, Übung.</p></body></html>",
        encoding="utf-8",
    )
    (HERE / "sample.csv").write_text("Position,Betrag\nÄpfel,12.5\n", encoding="utf-8")
    (HERE / "unsupported.xyz").write_bytes(b"\x00\x01binary")


if __name__ == "__main__":
    main()
